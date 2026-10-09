"""Excel import of the item master (Milestone 2).

One sheet, one row per item. Existing items (same name) are updated, so a corrected sheet can
be imported again. The whole file is rejected if any row has an error, so nothing half-saves.
"""

import io
from decimal import Decimal, InvalidOperation

from openpyxl import Workbook, load_workbook
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.core.errors import BusinessRuleError
from app.core.tenancy import TENANT_ID
from app.models.masters import Item
from app.schemas.items import ImportResult, ImportRowError, ItemCreate, ItemUnitIn, ItemUpdate
from app.services import items as item_service

COLUMNS = [
    "name",
    "category",
    "brand",
    "hsn",
    "gst_rate",
    "base_unit",
    "base_whole_only",
    "size",
    "grade",
    "weight_per_piece_kg",
    "min_margin",
    "units",
]
REQUIRED = {"name", "category", "hsn", "gst_rate", "base_unit"}
MAX_ROWS = 2000
EXAMPLE = [
    "TMT bar 12 mm Fe500D",
    "tmt",
    "Kamachi",
    "72142090",
    "18",
    "kg",
    "no",
    "12 mm",
    "Fe500D",
    "10.656",
    "1.25",
    "ton:1000, piece:10.656:whole",
]


def template_bytes() -> bytes:
    wb = Workbook()
    ws = wb.active
    if ws is None:  # pragma: no cover - a new workbook always has an active sheet
        raise RuntimeError("workbook has no sheet")
    ws.title = "Items"
    ws.append(COLUMNS)
    ws.append(EXAMPLE)
    note = wb.create_sheet("How to fill")
    for line in (
        "One row per item. Re-importing a sheet updates items with the same name.",
        "category: tmt, pipe, cement, wire, angle, channel or other.",
        "hsn: 4 to 8 digits. gst_rate: percent, e.g. 18.",
        "base_unit: what stock is counted in (kg, bag, piece).",
        "base_whole_only: yes for bags and pieces that cannot be split, otherwise no.",
        "units: other units, as unit:factor_to_base[:whole], separated by commas.",
        "  ton:1000 means 1 ton = 1000 base units. Add :whole for bags and pieces.",
        "min_margin: owner-only warning margin per base unit. Leave blank for 0.",
    ):
        note.append([line])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _cell(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _parse_units(text: str | None) -> list[ItemUnitIn]:
    units: list[ItemUnitIn] = []
    for part in (text or "").split(","):
        part = part.strip()
        if not part:
            continue
        pieces = [p.strip() for p in part.split(":")]
        if len(pieces) not in (2, 3) or (len(pieces) == 3 and pieces[2].lower() != "whole"):
            raise ValueError(f"Unit {part!r} must look like ton:1000 or bag:1:whole")
        try:
            factor = Decimal(pieces[1])
        except InvalidOperation as exc:
            raise ValueError(f"Unit {part!r}: factor {pieces[1]!r} is not a number") from exc
        units.append(ItemUnitIn(unit=pieces[0], factor_to_base=factor, whole_only=len(pieces) == 3))
    return units


def _row_to_create(row: dict[str, str | None]) -> ItemCreate:
    payload: dict[str, object] = {k: v for k, v in row.items() if k != "units" and v is not None}
    for key in ("hsn", "gst_rate", "weight_per_piece_kg", "min_margin"):
        if key in payload:
            payload[key] = str(payload[key])
    if "base_whole_only" in payload:
        flag = str(payload["base_whole_only"]).strip().lower()
        if flag not in {"yes", "no", "true", "false", "1", "0", "y", "n"}:
            raise ValueError("base_whole_only must be yes or no")
        payload["base_whole_only"] = flag in {"yes", "true", "1", "y"}
    payload["units"] = _parse_units(row.get("units"))
    return ItemCreate.model_validate(payload)


def import_items(db: Session, content: bytes, *, dry_run: bool) -> ImportResult:
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several types for bad files
        raise BusinessRuleError(
            "This is not a valid .xlsx file. Download the template and fill it in.",
            code="BAD_SPREADSHEET",
        ) from exc
    ws = wb["Items"] if "Items" in wb.sheetnames else wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip().lower() if c is not None else "" for c in next(rows, ())]
    missing = REQUIRED - set(header)
    if missing:
        raise BusinessRuleError(
            f"Missing columns: {', '.join(sorted(missing))}. Use the template.",
            code="MISSING_COLUMNS",
        )

    errors: list[ImportRowError] = []
    parsed: list[tuple[int, ItemCreate]] = []
    seen: dict[str, int] = {}
    for number, values in enumerate(rows, start=2):
        if number - 1 > MAX_ROWS:
            errors.append(
                ImportRowError(row=number, field=None, message=f"At most {MAX_ROWS} rows")
            )
            break
        row = {h: _cell(v) for h, v in zip(header, values, strict=False) if h in COLUMNS}
        if not any(row.values()):
            continue
        try:
            data = _row_to_create(row)
        except ValidationError as exc:
            for err in exc.errors():
                loc = ".".join(str(p) for p in err["loc"])
                errors.append(ImportRowError(row=number, field=loc or None, message=err["msg"]))
            continue
        except ValueError as exc:
            errors.append(ImportRowError(row=number, field="units", message=str(exc)))
            continue
        key = data.name.lower()
        if key in seen:
            errors.append(
                ImportRowError(row=number, field="name", message=f"Same name as row {seen[key]}")
            )
            continue
        seen[key] = number
        parsed.append((number, data))

    if errors:
        return ImportResult(dry_run=dry_run, created=0, updated=0, errors=errors)

    existing = {
        i.name.lower(): i
        for i in db.execute(
            select(Item).where(Item.tenant_id == TENANT_ID, func.lower(Item.name).in_(list(seen)))
        ).scalars()
    }
    created = updated = 0
    for number, data in parsed:
        current = existing.get(data.name.lower())
        if current is None:
            created += 1
            if not dry_run:
                item_service.apply_create(db, data)
        else:
            updated += 1
            if current.base_unit != data.base_unit:
                errors.append(
                    ImportRowError(
                        row=number,
                        field="base_unit",
                        message=f"{current.name!r} is counted in {current.base_unit}; "
                        "the base unit cannot change",
                    )
                )
            elif not dry_run:
                item_service.apply_update(
                    db, current, ItemUpdate.model_validate(data.model_dump(exclude={"base_unit"}))
                )
    if errors:
        db.rollback()
        return ImportResult(dry_run=dry_run, created=0, updated=0, errors=errors)
    if dry_run:
        db.rollback()
    else:
        db.commit()
    return ImportResult(dry_run=dry_run, created=created, updated=updated, errors=[])
