"""Milestone 13: GSTR-1 tables, GSTR-3B figures and GSTR-2B matching, with hand-worked numbers."""

import io
import json

import pytest
from openpyxl import load_workbook

from app.core.clock import today_ist
from app.domain.gst import gstin_checksum
from tests.conftest import login
from tests.integration.test_sales import (  # noqa: F401 - the world fixture is shared
    ACCOUNTANT,
    COUNTER,
    body,
    line,
    post,
    world,
)

pytestmark = pytest.mark.integration

SUPPLIER_GSTIN = "33BBBBB1234B1Z" + gstin_checksum("33BBBBB1234B1Z")
PERIOD = today_ist().strftime("%Y-%m")


@pytest.fixture
def books(client, world, seeded):
    """A month of mixed bills: B2B, a large inter-state B2C, small B2C, returns and a purchase."""
    owner = world["owner"]
    client.patch(
        f"/api/v1/parties/{world['supplier']['id']}", headers=owner, json={"gstin": SUPPLIER_GSTIN}
    )
    far = client.post(
        "/api/v1/parties",
        headers=owner,
        json={
            "name": "Karnataka buyer",
            "type": "customer",
            "state_code": "29",
            "sites": [{"name": "Hosur yard", "state_code": "29", "address": "Hosur"}],
        },
    ).json()
    client.patch(
        f"/api/v1/parties/{far['id']}",
        headers=owner,
        json={"credit_allowed": True, "credit_limit": "10000000", "credit_days": 30},
    )
    world["far"] = far
    world["far_site"] = far["sites"][0]["id"]
    a = post(
        client, world, body(world, [line(world, "tmt", "2.5", "ton"), line(world, "cement", "7")])
    ).json()
    c = post(
        client,
        world,
        {
            "location_id": world["loc"]["S1"],
            "party_id": far["id"],
            "site_id": world["far_site"],
            "lines": [line(world, "tmt", "2", "ton")],
        },
    ).json()
    d = post(client, world, body(world, [line(world, "cement", "7")], party="walkin")).json()
    assert c["supply_kind"] == "inter_state" and c["grand_total"] == "132160.00"
    cn1 = client.post(
        "/api/v1/credit-notes",
        headers=owner,
        json={
            "invoice_id": a["id"],
            "reason": "Too much",
            "lines": [{"line_id": a["lines"][0]["id"], "quantity": "1"}],
        },
    ).json()
    cn2 = client.post(
        "/api/v1/credit-notes",
        headers=owner,
        json={
            "invoice_id": d["id"],
            "reason": "Extra bags",
            "lines": [{"line_id": d["lines"][0]["id"], "quantity": "3"}],
        },
    ).json()
    purchase = client.post(
        "/api/v1/purchases",
        headers=owner,
        json={
            "supplier_id": world["supplier"]["id"],
            "location_id": world["loc"]["S1"],
            "bill_no": "B-1",
            "bill_date": today_ist().isoformat(),
            "lines": [
                {"item_id": world["tmt"]["id"], "unit": "ton", "quantity": "10", "rate": "55000"}
            ],
        },
    ).json()
    client.post(
        "/api/v1/debit-notes",
        headers=owner,
        json={
            "purchase_id": purchase["id"],
            "reason": "Rusted",
            "lines": [{"line_id": purchase["lines"][0]["id"], "quantity": "2"}],
        },
    )
    return {"a": a, "c": c, "d": d, "cn1": cn1, "cn2": cn2, "purchase": purchase}


def get(client, headers, path, **params):
    query = "&".join(f"{k}={v}" for k, v in params.items())
    return client.get(
        f"/api/v1/gst/{path}?period={PERIOD}" + (f"&{query}" if query else ""), headers=headers
    )


def test_gstr1_tables(client, world, books):
    r = get(client, world["owner"], "gstr1").json()
    (b2b,) = r["b2b"]
    assert (b2b["ctin"], b2b["number"], b2b["taxable"], b2b["cgst"], b2b["pos"]) == (
        "33AAPFU0939F1Z2",
        books["a"]["number"],
        "142663.15",
        "12972.84",
        "33",
    )
    assert [x["rate"] for x in b2b["rates"]] == ["18.00", "28.00"]
    (b2cl,) = r["b2cl"]
    assert (b2cl["ctin"], b2cl["pos"], b2cl["taxable"], b2cl["igst"], b2cl["value"]) == (
        None,
        "29",
        "112000.00",
        "20160.00",
        "132160.00",
    )
    # The walk-in's 7 bags (2,663.15) less the 3 returned (1,141.35): 1,521.80, tax 213.05 each.
    (b2cs,) = r["b2cs"]
    assert (
        b2cs["supply"],
        b2cs["pos"],
        b2cs["rate"],
        b2cs["taxable"],
        b2cs["cgst"],
        b2cs["sgst"],
    ) == ("INTRA", "33", "28.00", "1521.80", "213.05", "213.05")
    (cdnr,) = r["cdnr"]
    assert (cdnr["ctin"], cdnr["invoice_number"], cdnr["taxable"], cdnr["cgst"]) == (
        "33AAPFU0939F1Z2",
        books["a"]["number"],
        "56000.00",
        "5040.00",
    )
    assert r["cdnur"] == []
    by_hsn = {h["hsn"]: h for h in r["hsn"]}
    tmt, cement = by_hsn["72142090"], by_hsn["25232930"]
    # TMT: 2,500 + 2,000 kg sold less 1,000 returned = 3,500 kg; taxable 1,40,000 + 1,12,000 - 56,000.
    assert (tmt["quantity"], tmt["taxable"], tmt["uqc"], tmt["igst"], tmt["cgst"]) == (
        "3500.000",
        "196000.00",
        "KGS",
        "20160.00",
        "7560.00",
    )
    assert (cement["quantity"], cement["taxable"], cement["uqc"]) == ("11.000", "4184.95", "BAG")
    invoices = next(d for d in r["docs"] if d["nature"].startswith("Invoices"))
    assert (invoices["count"], invoices["first"], invoices["last"], invoices["gaps"]) == (
        3,
        books["a"]["number"],
        books["d"]["number"],
        [],
    )
    notes = next(d for d in r["docs"] if d["nature"] == "Credit note")
    assert notes["count"] == 2
    assert (
        r["totals"]["invoices"],
        r["totals"]["notes"],
        r["totals"]["taxable"],
        r["totals"]["igst"],
        r["totals"]["cgst"],
    ) == (3, 2, "200184.95", "20160.00", "8145.89")


def test_exports_as_portal_json_and_excel(client, world, books):
    owner = world["owner"]
    out = get(client, owner, "gstr1/export", format="json")
    assert out.status_code == 200 and "GSTR1-" in out.headers["content-disposition"]
    doc = json.loads(out.content)
    year, month = today_ist().year, today_ist().month
    assert doc["fp"] == f"{month:02d}{year}"
    assert doc["b2b"][0]["ctin"] == "33AAPFU0939F1Z2"
    inv = doc["b2b"][0]["inv"][0]
    assert (inv["inum"], inv["pos"], inv["rchrg"], inv["inv_typ"]) == (
        books["a"]["number"],
        "33",
        "N",
        "R",
    )
    assert inv["idt"] == today_ist().strftime("%d-%m-%Y")
    assert [i["itm_det"]["rt"] for i in inv["itms"]] == [18.0, 28.0]
    assert doc["b2cl"][0]["pos"] == "29" and doc["b2cs"][0]["typ"] == "OE"
    assert (
        doc["cdnr"][0]["nt"][0]["ntty"] == "C"
        and doc["cdnr"][0]["nt"][0]["inum"] == books["a"]["number"]
    )
    assert doc["hsn"]["data"][0]["hsn_sc"] and doc["doc_issue"]["doc_det"][0]["doc_num"] == 1
    sheet = get(client, owner, "gstr1/export", format="xlsx")
    assert sheet.status_code == 200
    book = load_workbook(io.BytesIO(sheet.content))
    assert book.sheetnames == [
        "Summary",
        "B2B",
        "B2CL",
        "B2CS",
        "CDNR",
        "CDNUR",
        "HSN",
        "Documents",
    ]
    assert book["B2B"]["C2"].value == books["a"]["number"]


def test_gstr3b_and_2b_matching(client, world, books):
    owner = world["owner"]
    r = get(client, owner, "gstr3b").json()
    assert (
        r["outward_taxable"]["taxable"],
        r["outward_taxable"]["igst"],
        r["outward_taxable"]["cgst"],
    ) == ("200184.95", "20160.00", "8145.89")
    assert r["outward_nil"] == "0.00" and r["itc_in_2b"] is None
    # 10 ton at ₹55,000 = ₹5,50,000: input CGST and SGST 49,500 each; the 2-ton return takes 9,900 back.
    assert (r["itc_books"]["cgst"], r["itc_books"]["sgst"]) == ("49500.00", "49500.00")
    assert (r["itc_reversed"]["taxable"], r["itc_reversed"]["cgst"]) == ("110000.00", "9900.00")
    assert r["net_payable"]["cgst"] == "-31454.11" and r["net_payable"]["igst"] == "20160.00"

    nothing = get(client, owner, "gstr2b").json()
    assert nothing["imported_rows"] == 0 and nothing["counts"]["missing_in_2b"] == 1
    csv_file = (
        "gstin,supplier,invoice_no,invoice_date,taxable,igst,cgst,sgst\n"
        f"{SUPPLIER_GSTIN},Steel Mills,b 1,01-10-2026,550000,0,49500,49500\n"
        "29ABCDE1234F1Z5,Other mill,Z-77,02-10-2026,1000,0,90,90\n"
    )
    up = client.post(
        "/api/v1/gst/gstr2b",
        headers=owner,
        data={"period": PERIOD},
        files={"file": ("2b.csv", csv_file, "text/csv")},
    )
    assert up.status_code == 201 and up.json()["row_count"] == 2
    m = get(client, owner, "gstr2b").json()
    assert m["counts"] == {"mismatch": 0, "missing_in_2b": 0, "missing_in_books": 1, "matched": 1}
    matched = next(x for x in m["rows"] if x["status"] == "matched")
    assert (matched["supplier"], matched["books_taxable"], matched["portal_tax"]) == (
        "Mills",
        "550000.00",
        "99000.00",
    )
    assert get(client, owner, "gstr3b").json()["itc_in_2b"]["cgst"] == "49590.00"

    # A newer upload replaces the old one: this time the supplier's tax is ₹20 short.
    short = csv_file.replace("49500,49500", "49490,49490")
    client.post(
        "/api/v1/gst/gstr2b",
        headers=owner,
        data={"period": PERIOD},
        files={"file": ("2b-v2.csv", short, "text/csv")},
    )
    again = get(client, owner, "gstr2b").json()
    assert again["file_name"] == "2b-v2.csv" and again["counts"]["mismatch"] == 1
    assert (
        again["rows"][0]["status"] == "mismatch" and again["rows"][0]["difference_tax"] == "-20.00"
    )


def test_2b_in_the_portal_json_layout(client, world, books):
    doc = {
        "data": {
            "docdata": {
                "b2b": [
                    {
                        "ctin": SUPPLIER_GSTIN,
                        "trdnm": "Steel Mills",
                        "inv": [
                            {
                                "inum": "B-1",
                                "dt": "09-10-2026",
                                "val": 649000,
                                "items": [
                                    {
                                        "num": 1,
                                        "rt": 18,
                                        "txval": 550000,
                                        "cgst": 49500,
                                        "sgst": 49500,
                                        "igst": 0,
                                    }
                                ],
                            }
                        ],
                    }
                ]
            }
        }
    }
    up = client.post(
        "/api/v1/gst/gstr2b",
        headers=world["owner"],
        data={"period": PERIOD},
        files={"file": ("gstr2b.json", json.dumps(doc), "application/json")},
    )
    assert up.status_code == 201, up.text
    assert get(client, world["owner"], "gstr2b").json()["counts"]["matched"] == 1


def test_bad_inputs_and_roles(client, world, books):
    owner, accountant, counter = world["owner"], login(client, *ACCOUNTANT), login(client, *COUNTER)
    for bad in ("2026-13", "oct", ""):
        r = client.get(f"/api/v1/gst/gstr1?period={bad}", headers=owner)
        assert r.status_code in (409, 422)
    assert (
        client.get("/api/v1/gst/gstr1?period=2026-13", headers=owner).json()["code"] == "BAD_PERIOD"
    )
    garbage = client.post(
        "/api/v1/gst/gstr2b",
        headers=owner,
        data={"period": PERIOD},
        files={"file": ("x.csv", "hello,world\n1,2\n", "text/csv")},
    )
    assert garbage.status_code == 400 and garbage.json()["code"] == "GSTR2B_UNREADABLE"
    empty = client.post(
        "/api/v1/gst/gstr2b",
        headers=owner,
        data={"period": PERIOD},
        files={
            "file": (
                "x.csv",
                "gstin,supplier,invoice_no,invoice_date,taxable,igst,cgst,sgst\n",
                "text/csv",
            )
        },
    )
    assert empty.json()["code"] == "GSTR2B_EMPTY"
    bad_json = client.post(
        "/api/v1/gst/gstr2b",
        headers=owner,
        data={"period": PERIOD},
        files={"file": ("x.json", "{not json", "application/json")},
    )
    assert bad_json.json()["code"] == "GSTR2B_UNREADABLE"
    wrong_number = client.post(
        "/api/v1/gst/gstr2b",
        headers=owner,
        data={"period": PERIOD},
        files={
            "file": (
                "x.csv",
                "gstin,supplier,invoice_no,invoice_date,taxable\nA,B,C,D,not-a-number\n",
                "text/csv",
            )
        },
    )
    assert wrong_number.json()["code"] == "GSTR2B_UNREADABLE"
    assert (
        client.post(
            "/api/v1/gst/gstr2b",
            headers=owner,
            data={"period": "bad"},
            files={"file": ("x.csv", "a", "text/csv")},
        ).json()["code"]
        == "BAD_PERIOD"
    )

    for path in ("gstr1", "gstr1/export", "gstr3b", "gstr2b"):
        assert get(client, accountant, path).status_code == 200
        assert get(client, counter, path).status_code == 403
    up = {
        "data": {"period": PERIOD},
        "files": {
            "file": (
                "2b.csv",
                "gstin,supplier,invoice_no,invoice_date,taxable,igst,cgst,sgst\nA,B,C,D,1,0,0,0\n",
                "text/csv",
            )
        },
    }
    assert client.post("/api/v1/gst/gstr2b", headers=counter, **up).status_code == 403
    assert client.post("/api/v1/gst/gstr2b", headers=accountant, **up).status_code == 201


def test_an_empty_month_and_prior_months(client, world, books):
    r = client.get("/api/v1/gst/gstr1?period=2020-01", headers=world["owner"]).json()
    assert r["b2b"] == [] and r["totals"]["invoices"] == 0 and r["totals"]["taxable"] == "0.00"
    assert (
        client.get("/api/v1/gst/gstr3b?period=2020-01", headers=world["owner"]).json()[
            "outward_nil"
        ]
        == "0.00"
    )
